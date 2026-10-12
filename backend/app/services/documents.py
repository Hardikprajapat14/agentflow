"""Upload handling, text extraction and indexing for documents."""
import logging
import os
import re
import uuid

import pandas as pd
from docx import Document as DocxDocument
from fastapi import HTTPException, UploadFile
from pypdf import PdfReader
from sqlalchemy.orm import Session

from app.config import get_settings
from app.models import Document, User
from app.rag.store import index_text
from app.services.audit import audit
from app.utils.files import ensure_local_file

logger = logging.getLogger(__name__)
ALLOWED = {"pdf", "docx", "csv", "xlsx", "txt", "md"}


def save_upload(db: Session, user: User, file: UploadFile) -> Document:
    s = get_settings()
    name = os.path.basename(file.filename or "upload")
    name = re.sub(r"[^\w.\- ]", "_", name)[:200] or "upload"
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if ext not in ALLOWED:
        raise HTTPException(415, f"Unsupported file type '.{ext}'. Allowed: {', '.join(sorted(ALLOWED))}")
    folder = os.path.join(s.upload_dir, str(user.id))
    os.makedirs(folder, exist_ok=True)
    path = os.path.join(folder, f"{uuid.uuid4().hex}.{ext}")
    size, limit = 0, s.max_upload_mb * 1024 * 1024
    with open(path, "wb") as out:
        while chunk := file.file.read(1024 * 1024):
            size += len(chunk)
            if size > limit:
                out.close()
                os.remove(path)
                raise HTTPException(413, f"File exceeds the {s.max_upload_mb} MB limit")
            out.write(chunk)
    if size == 0:
        os.remove(path)
        raise HTTPException(400, "The uploaded file is empty")
    with open(path, "rb") as fh:
        content = fh.read()
    doc = Document(user_id=user.id, filename=name, file_type=ext, file_size=size, storage_path=path,
                   status="processing", content=content)
    db.add(doc)
    db.commit()
    audit(db, "document_uploaded", f"Uploaded {name}", user_id=user.id)
    return doc


def extract_text(doc: Document) -> tuple[str, dict]:
    p, t = ensure_local_file(doc), doc.file_type
    if t == "pdf":
        reader = PdfReader(p)
        pages = [(pg.extract_text() or "") for pg in reader.pages]
        return "\n\n".join(pages), {"pages": len(pages)}
    if t == "docx":
        d = DocxDocument(p)
        parts = [para.text for para in d.paragraphs if para.text.strip()]
        for table in d.tables:
            parts += [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        return "\n\n".join(parts), {}
    if t in ("csv", "xlsx"):
        df = pd.read_csv(p) if t == "csv" else pd.read_excel(p)
        df.columns = [str(c).strip() for c in df.columns]
        text = (f"Dataset {doc.filename}: {len(df)} rows. Columns: {', '.join(df.columns)}.\n\n"
                f"Column statistics:\n{df.describe(include='all').round(2).to_string()[:3000]}\n\n"
                f"Sample rows:\n{df.head(15).to_string()}")
        return text, {"rows": len(df), "columns": list(df.columns)}
    with open(p, encoding="utf-8", errors="replace") as fh:
        return fh.read(), {}


def process_document(db: Session, document_id: int) -> None:
    doc = db.get(Document, document_id)
    if not doc:
        return
    try:
        text, meta = extract_text(doc)
        if not text.strip():
            raise ValueError("No extractable text found (is this a scanned PDF?)")
        n = index_text(db, doc, text)
        doc.meta, doc.status, doc.error = {**meta, "chunks": n}, "ready", None
    except Exception as exc:  # noqa: BLE001
        logger.exception("Document %s processing failed", document_id)
        db.rollback()
        doc = db.get(Document, document_id)
        doc.status, doc.error = "failed", f"{type(exc).__name__}: {exc}"[:500]
    db.commit()
